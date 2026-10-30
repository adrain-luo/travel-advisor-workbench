"""Synthetic fixtures validate local domain rules, never live WeChat or AI compatibility."""
from __future__ import annotations

import copy
import json
import tempfile
import threading
import time
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

from backend.extraction import (
    _request, analyze_material, check_provider, extraction_cache_key, extract_file, parse_requirements,
)
from backend.exports import offer_lines, price_text, recommendation_text, render_comparison, source_evidence_lines
from backend.matching import match_products


def product_fixture(**changes):
    product = {"id": "fixture-product", "title": "测试夹具·云南六日游", "supplier": "夹具供应商", "destination": "云南", "departure_city": "上海", "duration_days": 6, "shopping": "none", "hotels": "四钻酒店", "source_item_ids": ["fixture-source"], "offers": [{"id": "fixture-offer", "price": 3000, "price_kind": "exact", "currency": "CNY", "unit": "per_person", "traveler": "adult", "occupancy": "双人间", "departure_date": "2026-11-01", "extras": [], "transport_included": True, "confirmed_at": None}]}
    product.update(changes)
    return product


def requirements_fixture(**changes):
    requirements = {"destination": "云南", "date_start": "2026-11-01", "date_end": "2026-11-03", "budget": 3500, "budget_unit": "per_person", "currency": "CNY", "shopping": "none", "traveler": "adult", "occupancy": "双人间", "transport_included": True}
    requirements.update(changes)
    return requirements


class MatchingTests(unittest.TestCase):
    def test_known_exact_offer_matches_without_mutating_inputs(self):
        product, req = product_fixture(), requirements_fixture()
        originals = copy.deepcopy((product, req))
        result = match_products([product], req)
        self.assertEqual(len(result["matched"]), 1)
        self.assertEqual((product, req), originals)

    def test_unknown_shopping_is_not_satisfied(self):
        result = match_products([product_fixture(shopping=None)], requirements_fixture())
        self.assertFalse(result["matched"])
        self.assertEqual(len(result["needs_verification"]), 1)

    def test_unknown_extras_and_from_price_are_not_budget_passes(self):
        product = product_fixture()
        product["offers"][0].update(extras=None, price_kind="from")
        result = match_products([product], requirements_fixture())
        self.assertFalse(result["matched"])
        self.assertGreaterEqual(len(result["needs_verification"][0]["issues"]), 2)

    def test_explicit_zero_extras_is_different_from_unknown(self):
        product = product_fixture()
        product["offers"][0]["extras"] = [{"name": "燃油费", "required": True, "amount": 0}]
        self.assertEqual(len(match_products([product], requirements_fixture())["matched"]), 1)
        product["offers"][0]["extras"][0]["amount"] = None
        self.assertFalse(match_products([product], requirements_fixture())["matched"])

    def test_mandatory_extra_breaks_budget(self):
        product = product_fixture()
        product["offers"][0]["extras"] = [{"name": "燃油费", "required": True, "amount": 600}]
        self.assertEqual(len(match_products([product], requirements_fixture())["excluded"]), 1)

    def test_no_cross_offer_condition_mixing(self):
        product = product_fixture()
        expensive = copy.deepcopy(product["offers"][0])
        expensive["price"] = 4000
        cheap_wrong_date = copy.deepcopy(product["offers"][0])
        cheap_wrong_date["departure_date"] = "2026-12-01"
        product["offers"] = [expensive, cheap_wrong_date]
        result = match_products([product], requirements_fixture())
        self.assertFalse(result["matched"])
        self.assertEqual(len(result["excluded"]), 1)

    def test_validity_interval_is_not_guaranteed_departure(self):
        product = product_fixture()
        product["offers"][0].update(departure_date=None, valid_from="2026-11-01", valid_to="2026-11-30")
        self.assertEqual(len(match_products([product], requirements_fixture())["needs_verification"]), 1)

    def test_boundary_date_inclusive_and_foreign_currency_unresolved(self):
        self.assertEqual(len(match_products([product_fixture()], requirements_fixture(date_end="2026-11-01"))["matched"]), 1)
        product = product_fixture()
        product["offers"][0]["currency"] = "USD"
        self.assertFalse(match_products([product], requirements_fixture())["matched"])

    def test_frontend_aliases_and_preferences(self):
        result = match_products([product_fixture()], {"budget_max": 3500, "budget_unit": "per_person", "shopping": "any", "hotel_preferences": "五星酒店", "preferences": ["hotels"]})
        self.assertTrue(result["matched"])

    def test_suppliers_with_same_title_remain_separate(self):
        products = [product_fixture(), product_fixture(id="second", supplier="另一供应商")]
        self.assertEqual(len(match_products(products, requirements_fixture())["matched"]), 2)

    def test_invalid_inputs_never_satisfy_constraints(self):
        self.assertFalse(match_products([product_fixture()], requirements_fixture(budget="待确认"))["matched"])
        self.assertFalse(match_products([product_fixture()], requirements_fixture(date_start="not a date"))["matched"])
        product = product_fixture()
        product["offers"][0]["extras"] = [{"name": "税费", "required": True, "amount": "待确认"}]
        self.assertEqual(len(match_products([product], requirements_fixture())["needs_verification"]), 1)

    def test_explicit_hard_constraint_checkboxes_are_authoritative(self):
        req = {"destination": "", "departure_city": None, "hard_constraints": ["destination", "departure_city"]}
        self.assertEqual(len(match_products([product_fixture()], req)["matched"]), 1)
        req = {"destination": "海南", "shopping": "none", "hard_constraints": ["shopping"]}
        self.assertEqual(len(match_products([product_fixture()], req)["matched"]), 1)

    def test_list_strings_are_delimited_not_characters(self):
        req = {"shopping": "none", "budget": None, "hard_constraints": "shopping,budget", "preferences": "云南，轻松", "exclusions": "购物店；强制自费"}
        result = match_products([product_fixture()], req)
        issues = result["needs_verification"][0]["issues"]
        self.assertEqual(len(issues), 2)
        self.assertIn("购物店", issues[0])


class ExtractionTests(unittest.TestCase):
    def test_text_encoding_csv_locations_and_no_publication_date(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.csv"
            path.write_text("目的地,价格\n云南,3000", encoding="utf-8-sig")
            result = extract_file(path)
            self.assertEqual(result["state"], "EXTRACTED")
            self.assertEqual(result["segments"][1]["location"], "row:2")
            self.assertNotIn("published_at", result)

    def test_content_addressed_document_and_image_paths(self):
        from docx import Document
        from openpyxl import Workbook
        from PIL import Image
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sha256blob"
            document = Document()
            document.add_paragraph("文档测试夹具")
            document.save(path)
            self.assertIn("文档测试夹具", extract_file(path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")["text"])
            book = Workbook()
            book.active.append(["云南", 3000])
            # openpyxl accepts binary streams even when the blob lacks an extension.
            book.save(path)
            result = extract_file(path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.assertIn("云南", result["text"])
            path.write_text("目的地,价格\n云南,3000", encoding="utf-8")
            self.assertEqual(extract_file(path, "text/csv")["segments"][1]["location"], "row:2")
            Image.new("RGB", (20, 20)).save(path, format="PNG")
            fixture = {"products": [], "issues": []}
            with patch("backend.extraction._request", return_value={"content": json.dumps(fixture), "model": "mock-only", "usage": None}) as request:
                result = analyze_material({"id": "fixture", "attachments": [{"path": str(path), "mime": "image/png", "name": "素材.png"}]}, {"base_url": "https://fixture.invalid/v1", "model": "text", "vision_model": "vision"})
            self.assertEqual(result["status"], "ANALYZED")
            self.assertTrue(request.call_args.kwargs["vision"])
            self.assertEqual(len(request.call_args.args[2][1]["content"]), 3)

    def test_scanned_pdf_and_docx_embedded_images_reach_vision(self):
        from docx import Document
        from PIL import Image
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            image_path = root / "fixture.png"
            Image.new("RGB", (60, 60), "red").save(image_path)
            pdf_path = root / "scan"
            Image.new("RGB", (60, 60), "red").save(pdf_path, format="PDF")
            docx_path = root / "document"
            document = Document()
            document.add_picture(str(image_path))
            document.save(docx_path)
            attachments = [{"path": str(pdf_path), "mime": "application/pdf", "name": "scan.pdf"}, {"path": str(docx_path), "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "name": "document.docx"}]
            with patch("backend.extraction._request", return_value={"content": '{"products":[],"issues":[]}', "model": "fixture", "usage": None}) as request:
                result = analyze_material({"id": "fixture", "attachments": attachments}, {"base_url": "https://fixture.invalid/v1", "model": "text", "vision_model": "vision"})
            self.assertEqual(result["status"], "PARTIAL")
            self.assertEqual(len(result["image_locations"]), 2)
            self.assertIn("page:1", result["image_locations"][0])
            self.assertTrue(request.call_args.kwargs["vision"])

    def test_image_is_waiting_for_vision_and_executable_is_unsupported(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as directory:
            image = Path(directory) / "fixture.png"
            Image.new("RGB", (20, 20)).save(image)
            self.assertEqual(extract_file(image)["state"], "NEEDS_VISION")
            executable = Path(directory) / "fixture.exe"
            executable.write_bytes(b"not executed")
            self.assertEqual(extract_file(executable)["state"], "UNSUPPORTED")

    def test_offline_has_no_fabricated_products(self):
        result = analyze_material({"id": "fixture", "text": "云南成人 3000 元起"}, {}, None)
        self.assertEqual(result["status"], "NEEDS_CONFIGURATION")
        self.assertEqual(result["products"], [])

    def test_rules_only_does_not_claim_general_nlp(self):
        result = parse_requirements("六个人去云南，预算每人3000元，6天无购物", {}, None)
        self.assertEqual(result["status"], "RULES_ONLY")
        self.assertIsNone(result["destination"])
        self.assertEqual(result["shopping"], "none")

    def test_multi_product_extraction_and_invalid_evidence(self):
        fixture = {"products": [{"title": "产品甲", "supplier": "甲供应商", "offers": [{"price": 1000, "price_kind": "from", "confirmed_at": "2026-01-01"}], "evidence": [{"field": "title", "source_item_id": "fixture", "quote": "产品甲"}]}, {"title": "产品乙", "supplier": "乙供应商", "evidence": [{"field": "price", "source_item_id": "wrong", "quote": "不存在"}]}], "issues": []}
        with patch("backend.extraction._request", return_value={"content": json.dumps(fixture), "model": "mock-only", "usage": {"total_tokens": 10}}):
            result = analyze_material({"id": "fixture", "text": "产品甲；产品乙"}, {"base_url": "https://fixture.invalid/v1", "model": "mock-only"})
        self.assertEqual(result["status"], "ANALYZED")
        self.assertEqual(len(result["products"]), 2)
        self.assertIsNone(result["products"][0]["offers"][0]["confirmed_at"])
        self.assertEqual(result["products"][1]["evidence"], [])

    def test_invalid_json_repair_is_bounded(self):
        with patch("backend.extraction._request", return_value={"content": "not JSON", "model": "mock-only", "usage": None}) as request:
            result = analyze_material({"id": "fixture", "text": "素材"}, {"base_url": "https://fixture.invalid/v1", "model": "mock-only"})
        self.assertEqual(request.call_count, 2)
        self.assertEqual(result["status"], "INVALID_RESPONSE")

    def test_http_timeout_and_rate_limit_are_bounded_and_actionable(self):
        import httpx
        cfg = {"base_url": "https://fixture.invalid/v1", "model": "mock-only", "timeout_seconds": 5}
        with patch("httpx.Client") as client, patch("backend.extraction.time.sleep"):
            post = client.return_value.__enter__.return_value.post
            post.side_effect = httpx.ReadTimeout("fixture timeout with secret text that must not escape")
            with self.assertRaisesRegex(ValueError, "AI 请求超时"):
                _request(cfg, "fixture-secret", [{"role": "user", "content": "fixture"}])
            self.assertEqual(post.call_count, 3)
        with patch("httpx.Client") as client, patch("backend.extraction.time.sleep"):
            post = client.return_value.__enter__.return_value.post
            post.return_value = httpx.Response(429, json={"message": "fixture rate limit"})
            with self.assertRaisesRegex(ValueError, "AI 服务限流"):
                _request(cfg, None, [{"role": "user", "content": "fixture"}])
            self.assertEqual(post.call_count, 3)

    def test_provider_limit_applies_across_all_request_callers(self):
        guard = threading.Lock()
        active = 0
        peak = 0

        def boundary(*args, **kwargs):
            nonlocal active, peak
            with guard:
                active += 1
                peak = max(peak, active)
            time.sleep(0.03)
            with guard:
                active -= 1
            return {"content": "OK", "usage": None}

        settings = {"max_concurrency": 1, "timeout_seconds": 5}
        with patch("backend.extraction._request_unlimited", side_effect=boundary):
            with ThreadPoolExecutor(max_workers=4) as pool:
                outputs = list(pool.map(lambda _: _request(settings, None, []), range(4)))
        self.assertEqual(len(outputs), 4)
        self.assertEqual(peak, 1)

    def test_long_image_crops_preserve_positions_and_limits_are_explicit(self):
        from PIL import Image
        from backend.extraction import _image_payloads
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "poster.png"
            Image.new("RGB", (600, 5000), "white").save(path)
            images = _image_payloads({"path": str(path), "name": "poster.png", "mime": "image/png"})
            self.assertEqual(len(images), 3)
            self.assertIn("pixels:0,1880,600,3880", images[1]["location"])

    def test_cache_changes_with_model_and_material_not_secret(self):
        material = {"id": "fixture", "text": "原文"}
        key = extraction_cache_key(material, {"model": "one", "api_key": "secret-a"})
        self.assertEqual(key, extraction_cache_key(material, {"model": "one", "api_key": "secret-b"}))
        self.assertNotEqual(key, extraction_cache_key(material, {"model": "two"}))
        self.assertNotEqual(key, extraction_cache_key({**material, "text": "修订"}, {"model": "one"}))

    def test_provider_check_has_real_result_semantics(self):
        responses = [{"content": "OK", "usage": None}, {"content": '{"probe":"ok"}', "usage": None}, {"content": "blue", "usage": None}]
        with patch("backend.extraction._request", side_effect=responses):
            result = check_provider({"model": "fixture", "vision_model": "fixture"})
        self.assertEqual(result["text"], "PASS")
        self.assertEqual(result["structured"], "PASS")
        self.assertEqual(result["vision"], "PARTIAL")


class ExportTests(unittest.TestCase):
    def test_source_snapshots_preserve_relative_and_minute_evidence_separately(self):
        product = product_fixture(source_snapshots=[
            {"source_name": "所选朋友圈", "author": "发帖人", "source_time_raw": "昨天 09:07", "observed_at": "2026-10-25", "observation_time_precision": "day", "time_evidence": {"precision": "minute", "status": "parsed", "timezone": "Asia/Shanghai", "start": "2026-10-24T09:07:00+08:00"}},
            {"source_name": "所选微信群", "author": "群消息发送人", "source_time_raw": "2026年10月25日 21:55", "observed_at": "2026-10-25", "time_evidence": {"precision": "minute", "status": "parsed", "start": "2026-10-25T21:55:00+08:00"}},
            {"source_name": "未定时素材", "observed_at": "2026-10-25", "time_evidence": {"status": "unknown"}},
        ], updated_at="2026-10-26T00:45:00+08:00")
        lines = source_evidence_lines(product)
        self.assertEqual(len(lines), 3)
        self.assertIn("原文“昨天 09:07”", lines[0])
        self.assertIn("观察日期 2026-10-25", lines[0])
        self.assertIn("对应 2026-10-24 09:07", lines[0])
        self.assertIn("2026年10月25日 21:55", lines[1])
        self.assertNotIn("09:07:00", "\n".join(lines))
        self.assertNotIn("21:55:00", "\n".join(lines))
        self.assertIn("来源时间未知", lines[2])
        self.assertNotIn("00:45", "\n".join(lines))
        draft = recommendation_text([product], {"information_date": "2026-10-26"})
        self.assertIn(lines[0], draft)
        self.assertIn(lines[1], draft)

    def test_conditional_extra_retains_chinese_currency_unit_and_conditions(self):
        offer = product_fixture()["offers"][0]
        offer["extras"] = [{"name": "马鞍山起止附加费", "amount": 20, "currency": "CNY", "unit": "per_person", "required": None, "conditions": "仅马鞍山起止时适用"}]
        text = "\n".join(offer_lines(offer))
        self.assertIn("20 元/人", text)
        self.assertIn("是否必付待确认", text)
        self.assertIn("适用条件：仅马鞍山起止时适用", text)
        self.assertNotIn("CNY", text)

    def test_legacy_source_time_fallback_never_uses_ingestion_time(self):
        self.assertIn("来源时间未知", source_evidence_lines({"updated_at": "2026-10-26", "observed_at": "2026-10-25"})[0])
        self.assertIn("2026-10-24", source_evidence_lines({"source_time": "2026-10-24"})[0])

    def test_from_price_and_unknowns_remain_visible_in_chinese_draft(self):
        product = product_fixture()
        product["offers"][0].update(price_kind="from", extras=None)
        draft = recommendation_text([product], {"information_date": "2026-10-25"})
        self.assertIn("起", draft)
        self.assertIn("必付附加费用未知", draft)
        self.assertIn("尚未确认", draft)
        self.assertNotIn("已发送", draft)
        self.assertIn("2026-11-01", draft)

    def test_chinese_png_contains_exact_draft_without_clipping(self):
        from PIL import Image
        products = [product_fixture(), product_fixture(title="测试夹具·较长标题的另一个旅行产品，用于验证自动换行")]
        requirements = {"information_date": "2026-10-25"}
        draft = recommendation_text(products, requirements) + "\n顾问追加：请优先选择适合同行儿童的轻松行程。"
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture-comparison.png"
            render_comparison(products, requirements, draft, path)
            with Image.open(path) as image:
                self.assertEqual(image.width, 1440)
                self.assertGreater(image.height, 1000)
                self.assertEqual(image.format, "PNG")
            self.assertGreater(path.stat().st_size, 10000)


if __name__ == "__main__":
    unittest.main()
